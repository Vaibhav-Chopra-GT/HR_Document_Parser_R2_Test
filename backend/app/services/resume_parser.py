"""
Resume Parser Service - Extracts text from PDF/DOCX and uses AI for structured extraction
"""
import os
import json
from app.services.ai_service import get_ai_service


class ResumeParser:
    """Parse resumes and extract candidate information"""

    def __init__(self):
        self.ai_service = get_ai_service()

    def extract_text_from_pdf(self, file_path: str) -> str:
        """Extract text from PDF file"""
        import pdfplumber

        text_parts = []
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
        except Exception as e:
            raise ValueError(f"Failed to parse PDF: {str(e)}")

        return '\n'.join(text_parts)

    def extract_text_from_docx(self, file_path: str) -> str:
        """Extract text from DOCX file"""
        from docx import Document

        try:
            doc = Document(file_path)
            text_parts = []

            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)

            # Also extract from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_parts.append(' | '.join(row_text))

            return '\n'.join(text_parts)

        except Exception as e:
            raise ValueError(f"Failed to parse DOCX: {str(e)}")

    def extract_text(self, file_path: str) -> str:
        """Extract text from resume file (PDF or DOCX)"""
        ext = os.path.splitext(file_path)[1].lower()

        if ext == '.pdf':
            return self.extract_text_from_pdf(file_path)
        elif ext == '.docx':
            return self.extract_text_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

    def parse_resume(self, file_path: str) -> dict:
        """
        Parse resume and extract structured candidate information

        Returns:
            dict with extracted data, confidence scores, and raw text
        """
        # Step 1: Extract text
        try:
            raw_text = self.extract_text(file_path)
        except Exception as e:
            return {
                "success": False,
                "error": f"Text extraction failed: {str(e)}",
                "raw_text": None
            }

        if not raw_text or len(raw_text.strip()) < 50:
            return {
                "success": False,
                "error": "Resume appears to be empty or too short",
                "raw_text": raw_text
            }

        # Step 2: Use AI to extract structured data
        ai_result = self.ai_service.extract_resume_data(raw_text)

        if not ai_result.get("success"):
            return {
                "success": False,
                "error": f"AI extraction failed: {ai_result.get('error', 'Unknown error')}",
                "raw_text": raw_text
            }

        # Step 3: Process and validate extracted data
        extracted = ai_result.get("data", {})

        processed = {
            "success": True,
            "raw_text": raw_text,
            "extracted": {
                "name": self._get_field_value(extracted, "name"),
                "email": self._get_field_value(extracted, "email"),
                "phone": self._get_field_value(extracted, "phone"),
                "company": self._get_field_value(extracted, "company"),
                "designation": self._get_field_value(extracted, "designation"),
                "skills": self._get_field_value(extracted, "skills", default=[])
            },
            "confidence": {
                "name": self._get_confidence(extracted, "name"),
                "email": self._get_confidence(extracted, "email"),
                "phone": self._get_confidence(extracted, "phone"),
                "company": self._get_confidence(extracted, "company"),
                "designation": self._get_confidence(extracted, "designation"),
                "skills": self._get_confidence(extracted, "skills")
            }
        }

        # Calculate overall confidence
        confidences = [v for v in processed["confidence"].values() if v is not None]
        processed["confidence"]["overall"] = sum(confidences) / len(confidences) if confidences else 0.0

        return processed

    def _get_field_value(self, data: dict, field: str, default=None):
        """Extract field value from AI response"""
        field_data = data.get(field, {})
        if isinstance(field_data, dict):
            return field_data.get("value", default)
        return field_data if field_data else default

    def _get_confidence(self, data: dict, field: str) -> float:
        """Extract confidence score from AI response"""
        field_data = data.get(field, {})
        if isinstance(field_data, dict):
            conf = field_data.get("confidence", 0.0)
            # Ensure it's a valid float between 0 and 1
            try:
                conf = float(conf)
                return max(0.0, min(1.0, conf))
            except (TypeError, ValueError):
                return 0.0
        return 0.0
